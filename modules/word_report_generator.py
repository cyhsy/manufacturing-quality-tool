from datetime import date
from io import BytesIO

import plotly.io as pio
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor

from modules.defect_analysis import (
    PART_NAME_COLUMN,
    PROCESS_COLUMN,
    SEVERITY_COLUMN,
    SUPPLIER_COLUMN,
    format_percent,
    get_top_dimension_item,
    generate_quality_summary,
    summarize_by_date,
    summarize_by_defect_type,
    summarize_by_dimension,
)


def _add_placeholder(doc, text="【待填写】"):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.color.rgb = RGBColor(128, 128, 128)
    run.font.size = Pt(10)
    return p


def _fig_to_image(fig):
    """Convert a Plotly figure to PNG bytes for embedding in Word."""
    try:
        return pio.to_image(fig, format="png", width=800, height=450, scale=1.5)
    except Exception:
        return None


def _add_chart_image(doc, fig, caption):
    """Add a chart image to the document with a caption."""
    img_bytes = _fig_to_image(fig)
    if img_bytes:
        doc.add_picture(BytesIO(img_bytes), width=Inches(5.5))
        caption_p = doc.add_paragraph()
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption_p.add_run(caption)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(100, 100, 100)
    else:
        doc.add_paragraph(f"[图表：{caption} — 图片生成失败，请在应用中查看]")


def build_word_report(df, pareto_fig=None, trend_fig=None, histogram_fig=None, control_fig=None):
    """Build an 8D quality analysis report in Word format.

    Returns the document as bytes ready for download.
    """
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)
    style.font.name = "微软雅黑"

    title = doc.add_heading("8D 质量问题分析报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta_table = doc.add_table(rows=4, cols=2)
    meta_data = [
        ("报告编号", f"8D-{date.today().strftime('%Y%m%d')}"),
        ("报告日期", str(date.today())),
        ("数据来源", "质量异常记录"),
        ("分析工具", "制造质量数据分析与检验计划辅助工具"),
    ]
    for i, (label, value) in enumerate(meta_data):
        meta_table.cell(i, 0).text = label
        meta_table.cell(i, 1).text = value
    doc.add_paragraph()

    defect_summary = summarize_by_defect_type(df)
    total_defects = int(defect_summary["缺陷数量"].sum())
    top_defect = get_top_dimension_item(defect_summary)

    doc.add_heading("D1 - 改善团队", level=1)
    doc.add_paragraph("组建跨部门改善小组，明确各成员职责：")
    _add_placeholder(doc, "组长：【待指定】")
    _add_placeholder(doc, "质量工程师：【待指定】")
    _add_placeholder(doc, "制造工程师：【待指定】")
    _add_placeholder(doc, "工艺工程师：【待指定】")
    _add_placeholder(doc, "供应商代表：【视需要】")

    doc.add_heading("D2 - 问题描述", level=1)
    doc.add_paragraph(
        f"本期质量异常记录共计 {len(df)} 条，"
        f"总缺陷数量 {total_defects} 件，"
        f"涉及 {len(defect_summary)} 类缺陷。"
    )
    if top_defect is not None:
        top_name = top_defect["缺陷类型"]
        top_count = int(top_defect["缺陷数量"])
        if total_defects > 0:
            top_ratio = top_count / total_defects
        else:
            top_ratio = 0
        doc.add_paragraph(
            f"主要缺陷类型为「{top_name}」，共 {top_count} 件，"
            f"占总缺陷数量的 {format_percent(top_ratio)}。"
        )

    doc.add_paragraph("问题详细描述：")
    _add_placeholder(doc, "【请描述问题发生的具体现象、发现时间、影响范围】")

    doc.add_heading("D3 - 围堵措施", level=1)
    doc.add_paragraph("立即采取的临时措施，防止问题继续流出：")
    _add_placeholder(doc, "1. 【例：隔离可疑批次产品】")
    _add_placeholder(doc, "2. 【例：对在制品进行100%检验】")
    _add_placeholder(doc, "3. 【例：通知供应商暂停发货】")
    _add_placeholder(doc, "围堵措施实施日期：【待填写】")
    _add_placeholder(doc, "围堵措施验证结果：【待填写】")

    doc.add_heading("D4 - 根因分析", level=1)
    doc.add_paragraph("通过4M1E（人机料法环）方法分析问题根本原因：")

    quality_summary = generate_quality_summary(df)
    for item in quality_summary:
        doc.add_paragraph(item, style="List Bullet")

    part_summary = summarize_by_dimension(df, PART_NAME_COLUMN)
    process_summary = summarize_by_dimension(df, PROCESS_COLUMN)
    supplier_summary = summarize_by_dimension(df, SUPPLIER_COLUMN)

    doc.add_paragraph()
    doc.add_paragraph("重点问题维度：")

    top_part = get_top_dimension_item(part_summary)
    if top_part is not None:
        doc.add_paragraph(
            f"零件：{top_part[PART_NAME_COLUMN]}，缺陷 {int(top_part['缺陷数量'])} 件"
        )

    top_process = get_top_dimension_item(process_summary)
    if top_process is not None:
        doc.add_paragraph(
            f"工序：{top_process[PROCESS_COLUMN]}，缺陷 {int(top_process['缺陷数量'])} 件"
        )

    top_supplier = get_top_dimension_item(supplier_summary)
    if top_supplier is not None:
        doc.add_paragraph(
            f"供应商：{top_supplier[SUPPLIER_COLUMN]}，缺陷 {int(top_supplier['缺陷数量'])} 件"
        )

    doc.add_paragraph()
    _add_placeholder(doc, "根本原因确认：【请根据根因分析结果和现场验证填写】")

    doc.add_heading("D5 - 永久纠正措施", level=1)
    doc.add_paragraph("针对确认的根本原因，制定并实施永久纠正措施：")
    _add_placeholder(doc, "措施1：【例：更新刀具更换频次，从每500件改为每300件】")
    _add_placeholder(doc, "措施2：【例：增加进厂来料抽检比例】")
    _add_placeholder(doc, "措施3：【例：优化夹具设计，增加定位销】")
    _add_placeholder(doc, "计划完成日期：【待填写】")
    _add_placeholder(doc, "责任人：【待填写】")

    doc.add_heading("D6 - 措施实施与验证", level=1)
    doc.add_paragraph("验证永久纠正措施的有效性：")
    _add_placeholder(doc, "验证方法：【例：连续生产3批次，抽样检验】")
    _add_placeholder(doc, "验证周期：【待填写】")
    _add_placeholder(doc, "验证结果：【待填写】")
    _add_placeholder(doc, "验证人：【待填写】")

    doc.add_heading("D7 - 预防再发生", level=1)
    doc.add_paragraph("系统性预防措施，防止同类问题再次发生：")
    _add_placeholder(doc, "1. 【例：更新PFMEA，将此失效模式纳入分析】")
    _add_placeholder(doc, "2. 【例：修订作业指导书，增加关键控制点】")
    _add_placeholder(doc, "3. 【例：增加过程控制图的监控频次】")
    _add_placeholder(doc, "横向推广范围：【是否有类似工序/零件需要同步改进】")

    doc.add_heading("D8 - 总结与关闭", level=1)

    severity_summary = summarize_by_dimension(df, SEVERITY_COLUMN)
    high_count = 0
    for _, row in severity_summary.iterrows():
        if row[SEVERITY_COLUMN] == "高":
            high_count = int(row["缺陷数量"])
            break

    doc.add_paragraph(
        f"本次问题分析共识别 {len(defect_summary)} 类缺陷，"
        f"其中高严重等级缺陷 {high_count} 件。"
    )
    if top_defect is not None and total_defects > 0:
        doc.add_paragraph(
            f"主要改善方向：降低「{top_defect['缺陷类型']}」的发生率，"
            f"目标将相关缺陷减少50%以上。"
        )
    doc.add_paragraph()
    _add_placeholder(doc, "项目关闭评审日期：【待填写】")
    _add_placeholder(doc, "评审结论：【待填写】")
    _add_placeholder(doc, "批准人：【待填写】")

    doc.add_heading("附录 A：缺陷类型统计", level=1)
    table = doc.add_table(
        rows=min(len(defect_summary), 15) + 1, cols=3, style="Light Grid Accent 1"
    )
    table.cell(0, 0).text = "缺陷类型"
    table.cell(0, 1).text = "缺陷数量"
    table.cell(0, 2).text = "占比"
    for i, (_, row) in enumerate(defect_summary.head(15).iterrows()):
        table.cell(i + 1, 0).text = str(row["缺陷类型"])
        table.cell(i + 1, 1).text = str(int(row["缺陷数量"]))
        table.cell(i + 1, 2).text = format_percent(row["占比"])

    doc.add_heading("附录 B：质量分析图表", level=1)

    if pareto_fig is not None:
        doc.add_heading("帕雷托图", level=2)
        _add_chart_image(doc, pareto_fig, "图1：缺陷类型帕雷托图")

    if trend_fig is not None:
        doc.add_heading("日期趋势图", level=2)
        _add_chart_image(doc, trend_fig, "图2：按日期缺陷趋势图")

    if histogram_fig is not None:
        doc.add_heading("直方图", level=2)
        _add_chart_image(doc, histogram_fig, "图3：缺陷数量分布直方图")

    if control_fig is not None:
        doc.add_heading("控制图", level=2)
        _add_chart_image(doc, control_fig, "图4：每日缺陷数量控制图（I-chart）")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("— 报告由制造质量数据分析工具自动生成，部分内容需人工补充 —")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(150, 150, 150)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()
